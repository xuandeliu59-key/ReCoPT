#!/usr/bin/env python3

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


RQ2_ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = RQ2_ROOT / "cache/embeddings"
SAMPLE_ID_SOURCE = RQ2_ROOT / "data/reference/original_test_2k.jsonl"
OUTPUT_DIR = RQ2_ROOT / "results/representation"

REFERENCE_FILE = INPUT_DIR / "codebert_reference_10k.npy"
REFERENCE_IDS_FILE = INPUT_DIR / "codebert_reference_10k_ids.npy"
ORIGINAL_FILE = INPUT_DIR / "codebert_Original.npy"

METHODS = {
    "Random_View": ("Random-View", INPUT_DIR / "codebert_Random_View.npy"),
    "LLM_View": ("LLM-View", INPUT_DIR / "codebert_LLM_View.npy"),
    "ContraBERT_View": (
        "ContraBERT-View",
        INPUT_DIR / "codebert_ContraBERT_View.npy",
    ),
    "ReCoPT_View": ("ReCoPT-View", INPUT_DIR / "codebert_ReCoPT_View.npy"),
}

RIDGE = 1e-6


def summarize(values: np.ndarray) -> dict[str, float | int]:
    return {
        "n": int(values.size),
        "mean": float(np.mean(values)),
        "std": float(np.std(values)),
        "median": float(np.median(values)),
        "q1": float(np.quantile(values, 0.25)),
        "q3": float(np.quantile(values, 0.75)),
        "min": float(np.min(values)),
        "max": float(np.max(values)),
    }


def build_workbook(
    euclidean: dict[str, np.ndarray],
    raw_scores: dict[str, np.ndarray],
    standardized_scores: dict[str, np.ndarray],
    reference_mean: float,
    reference_std: float,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    headers = [
        "Method",
        "N",
        "Euclidean Mean",
        "Euclidean Std",
        "Euclidean Median",
        "Euclidean Q1",
        "Euclidean Q3",
        "Raw Mahalanobis Mean",
        "Standardized Mahalanobis Mean",
        "Standardized Mahalanobis Std",
        "Standardized Mahalanobis Median",
        "Standardized Mahalanobis Q1",
        "Standardized Mahalanobis Q3",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rows: list[list[float | int | str]] = []
    for method_key, (label, _) in METHODS.items():
        e = summarize(euclidean[method_key])
        raw = summarize(raw_scores[method_key])
        standardized = summarize(standardized_scores[method_key])
        rows.append(
            [
                label,
                e["n"],
                e["mean"],
                e["std"],
                e["median"],
                e["q1"],
                e["q3"],
                raw["mean"],
                standardized["mean"],
                standardized["std"],
                standardized["median"],
                standardized["q1"],
                standardized["q3"],
            ]
        )
        sheet.append(rows[-1])

    for column in range(3, len(headers) + 1):
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=column).number_format = "0.0000"

    for column in (3, 9):
        values = [sheet.cell(row=row, column=column).value for row in range(2, sheet.max_row + 1)]
        maximum = max(values)
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row, column=column)
            if cell.value == maximum:
                cell.font = Font(bold=True, color="006100")
                cell.fill = PatternFill("solid", fgColor="C6EFCE")

    widths = [20, 9] + [24] * (len(headers) - 2)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"

    metadata = workbook.create_sheet("Method")
    metadata.append(["Item", "Value"])
    metadata.append(["Encoder", "microsoft/codebert-base"])
    metadata.append(["Representation", "Final-layer CLS, 768 dimensions"])
    metadata.append(["Reference count", 10000])
    metadata.append(["Evaluation count", 2000])
    metadata.append(["Euclidean", "||h_view-h_source||_2 without L2 normalization"])
    metadata.append(
        [
            "Raw Mahalanobis",
            "0.5*(h-mu)^T*(scatter+1e-6*I)^-1*(h-mu)",
        ]
    )
    metadata.append(["Standardization", "max(0, (raw-reference_mean)/reference_std)"])
    metadata.append(["Reference raw mean", reference_mean])
    metadata.append(["Reference raw std", reference_std])
    metadata.column_dimensions["A"].width = 25
    metadata.column_dimensions["B"].width = 80
    for cell in metadata[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    workbook.save(OUTPUT_DIR / "CodeBERT_Distance_Summary.xlsx")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    reference = np.load(REFERENCE_FILE).astype(np.float64)
    reference_ids = np.load(REFERENCE_IDS_FILE, allow_pickle=False)
    original = np.load(ORIGINAL_FILE).astype(np.float64)
    method_embeddings = {
        key: np.load(path).astype(np.float64)
        for key, (_, path) in METHODS.items()
    }

    if reference.shape != (10000, 768):
        raise ValueError(f"Unexpected reference shape: {reference.shape}")
    if reference_ids.shape != (10000,) or len(set(reference_ids.tolist())) != 10000:
        raise ValueError("Reference IDs are not 10,000 unique values")
    if original.shape != (2000, 768):
        raise ValueError(f"Unexpected original shape: {original.shape}")
    for key, values in method_embeddings.items():
        if values.shape != original.shape:
            raise ValueError(f"Unexpected shape for {key}: {values.shape}")
    for name, values in {"Reference": reference, "Original": original, **method_embeddings}.items():
        if not np.isfinite(values).all():
            raise ValueError(f"Non-finite embeddings in {name}")

    sample_ids = np.asarray(sorted(
        str(json.loads(line).get("submission_id", json.loads(line).get("id", "")))
        for line in SAMPLE_ID_SOURCE.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ))
    if sample_ids.shape != (2000,) or len(set(sample_ids.tolist())) != 2000:
        raise ValueError("Expected 2,000 unique aligned sample IDs")

    euclidean = {
        key: np.linalg.norm(values - original, axis=1)
        for key, values in method_embeddings.items()
    }

    center = np.mean(reference, axis=0)
    centered = reference - center
    scatter = centered.T @ centered
    eigenvalues, eigenvectors = np.linalg.eigh(scatter)
    inverse_eigenvalues = 1.0 / (eigenvalues + RIDGE)

    def raw_score(features: np.ndarray) -> np.ndarray:
        projected = (features - center) @ eigenvectors
        return 0.5 * np.sum(
            projected * projected * inverse_eigenvalues,
            axis=1,
        )

    reference_raw = raw_score(reference)
    reference_mean = float(np.mean(reference_raw))
    reference_std = float(np.std(reference_raw))
    if reference_std < 1e-12:
        raise ValueError("Reference Mahalanobis standard deviation is zero")

    def standardized(raw: np.ndarray) -> np.ndarray:
        return np.maximum(0.0, (raw - reference_mean) / reference_std)

    reference_standardized = standardized(reference_raw)
    original_raw = raw_score(original)
    original_standardized = standardized(original_raw)
    raw_scores = {key: raw_score(values) for key, values in method_embeddings.items()}
    standardized_scores = {
        key: standardized(values) for key, values in raw_scores.items()
    }

    distance_arrays: dict[str, np.ndarray] = {
        "sample_ids": sample_ids,
        "encoder": np.asarray("microsoft/codebert-base"),
    }
    for key, values in euclidean.items():
        distance_arrays[f"{key}__euclidean"] = values
    np.savez_compressed(OUTPUT_DIR / "distance_scores.npz", **distance_arrays)

    mahalanobis_arrays: dict[str, np.ndarray] = {
        "reference_ids": reference_ids,
        "sample_ids": sample_ids,
        "encoder": np.asarray("microsoft/codebert-base"),
        "ridge": np.asarray(RIDGE),
        "reference_raw_mean": np.asarray(reference_mean),
        "reference_raw_std": np.asarray(reference_std),
        "Reference__raw": reference_raw,
        "Reference__standardized": reference_standardized,
        "Original__raw": original_raw,
        "Original__standardized": original_standardized,
    }
    for key in METHODS:
        mahalanobis_arrays[f"{key}__raw"] = raw_scores[key]
        mahalanobis_arrays[f"{key}__standardized"] = standardized_scores[key]
        mahalanobis_arrays[f"{key}__raw_delta"] = raw_scores[key] - original_raw
        mahalanobis_arrays[f"{key}__standardized_delta"] = (
            standardized_scores[key] - original_standardized
        )
    np.savez_compressed(OUTPUT_DIR / "mahalanobis_scores.npz", **mahalanobis_arrays)

    build_workbook(
        euclidean,
        raw_scores,
        standardized_scores,
        reference_mean,
        reference_std,
    )

    print("CodeBERT distance analysis completed")
    for key, (label, _) in METHODS.items():
        print(
            f"{label}: euclidean={np.mean(euclidean[key]):.6f} "
            f"raw_mahalanobis={np.mean(raw_scores[key]):.8f} "
            f"standardized_mahalanobis={np.mean(standardized_scores[key]):.6f}"
        )
    print(f"output={OUTPUT_DIR}")


if __name__ == "__main__":
    main()
